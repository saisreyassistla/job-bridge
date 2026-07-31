"""Stage 1 — Data sources: load PDF, HTML, DOCX, and plain text files."""

from __future__ import annotations
import os
from pathlib import Path
from typing import Iterator


class RawDocument:
    def __init__(self, content: str, source: str, page: int = 0, mime: str = "text/plain"):
        self.content = content
        self.source = source   # file path or URL
        self.page = page
        self.mime = mime

    def __repr__(self):
        return f"RawDocument(source={self.source!r}, page={self.page}, chars={len(self.content)})"


def load_pdf(path: str) -> list[RawDocument]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextContainer

    docs = []
    for page_num, layout in enumerate(extract_pages(path), start=1):
        text = "".join(
            element.get_text()
            for element in layout
            if isinstance(element, LTTextContainer)
        )
        if text.strip():
            docs.append(RawDocument(text, source=path, page=page_num, mime="application/pdf"))
    return docs


def load_docx(path: str) -> list[RawDocument]:
    from docx import Document
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [RawDocument(text, source=path, mime="application/vnd.openxmlformats")]


def load_html(path: str) -> list[RawDocument]:
    from bs4 import BeautifulSoup
    with open(path, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    text = soup.get_text(separator="\n", strip=True)
    return [RawDocument(text, source=path, mime="text/html")]


def load_txt(path: str) -> list[RawDocument]:
    with open(path, encoding="utf-8") as f:
        return [RawDocument(f.read(), source=path, mime="text/plain")]


_LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
    ".txt": load_txt,
    ".md": load_txt,
}


def load_file(path: str) -> list[RawDocument]:
    ext = Path(path).suffix.lower()
    loader = _LOADERS.get(ext)
    if not loader:
        raise ValueError(f"Unsupported file type: {ext}")
    return loader(path)


def load_directory(directory: str) -> Iterator[RawDocument]:
    """Recursively load all supported documents from a directory."""
    for root, _, files in os.walk(directory):
        for fname in files:
            ext = Path(fname).suffix.lower()
            if ext in _LOADERS:
                path = os.path.join(root, fname)
                try:
                    yield from load_file(path)
                except Exception as e:
                    print(f"Warning: failed to load {path}: {e}")
